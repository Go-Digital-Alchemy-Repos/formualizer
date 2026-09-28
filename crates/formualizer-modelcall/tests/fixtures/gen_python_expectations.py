"""Regenerate the Python expectations the Lane B tests compare against (GOD-383).

Synthetic inputs only (no client workbook). Runs the parity repository's own
code, so the Rust port is checked against the behaviour source:

    python3 gen_python_expectations.py <parity checkout>

Writes, next to this file: prefetch_plan.xlsx (built with openpyxl),
prefetch_expected.json (workbook_runtime.prefetch.sibling_plan_for_path and
sibling_vectors), brent_expected.json (replayer.engine_adapter._brent_solve)
and goal_seek_expected.json (workbook_runtime.solve.run_solves over a fake
workbook).
"""
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(Path(sys.argv[1]).resolve()))

from openpyxl import Workbook  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402

from workbook_runtime import prefetch  # noqa: E402
from replayer import engine_adapter  # noqa: E402
from workbook_runtime import solve  # noqa: E402


def typed(value):
    if value is prefetch.DYNAMIC:
        return {'type': 'dynamic'}
    if value is prefetch.ABSENT:
        return {'type': 'absent'}
    if value is None:
        return {'type': 'none'}
    if isinstance(value, bool):
        return {'type': 'bool', 'value': value}
    if isinstance(value, int):
        return {'type': 'int', 'value': value}
    if isinstance(value, float):
        return {'type': 'float', 'value': repr(value)}
    if isinstance(value, str):
        return {'type': 'str', 'value': value}
    raise TypeError(type(value))


def untyped(item):
    kind, value = item[0], item[1]
    return {'int': int, 'float': float, 'str': str, 'bool': bool}[kind](value)


# -- sibling plan ----------------------------------------------------------------
X = '_xldudf_CS_SPARK_XCALL'
book = Workbook()
calc = book.active
calc.title = 'Calc'
inputs = book.create_sheet("My 'Inputs'")
data = book.create_sheet('Data')
data['B2'] = 99              # covered by Xinput_Rate: never a constant
inputs['B5'] = 0.05
inputs['B6'] = 7
inputs['B7'] = 'Gold'
inputs['B8'] = True
inputs['B9'] = '=1+1'        # a formula: dynamic
book.defined_names['Xinput_Rate'] = DefinedName('Xinput_Rate', attr_text='Data!$B$2')
book.defined_names['Xsolve_Goal'] = DefinedName('Xsolve_Goal', attr_text='Calc!$H$1:$I$5')
Q = "'My ''Inputs'''"
# Group 1: override (Term literal, Rate cell vs literal) and an extra pair.
calc['C2'] = f'={X}("Folder/Svc", $A$1:$B$4, "Premium", "Term", 10, "Rate", {Q}!B5)'
calc['C3'] = f'={X}("Folder/Svc", $A$1:$B$4, "Premium", "Term", 20, "Rate", {Q}!B5)'
calc['C4'] = f'={X}("Folder/Svc", $A$1:$B$4, "Premium", "Term", 30, "Rate", 0.06, "Plan", "Gold")'
# Group 2: partial (one shared dynamic argument) plus override.
calc['D2'] = f'=CS.SPARK.XCALL("Folder/Svc", A1:B4, "Value", "Age", D1, "Band", 1)'
calc['D3'] = f'=CS.SPARK.XCALL("Folder/Svc", A1:B4, "Value", "Age", D1, "Band", 2)'
calc['D4'] = f'=CS.SPARK.XCALL("Folder/Svc", A1:B4, "Value", "Age", 45, "Band", {Q}!B6)'
# Group 3: rejected (differing input blocks).
calc['E2'] = f'={X}("Folder/Svc", A1:B4, "Other", "Term", 1)'
calc['E3'] = f'={X}("Folder/Svc", A5:B8, "Other", "Term", 2)'
# Unresolved target, single-member group.
calc['F2'] = f'={X}(A9, A1:B4, "Premium")'
calc['F3'] = f'={X}("Lonely/Svc", 0, "Out")'
# Group 4: nested in IF, whitespace, negative and float literals, bool/text cells.
calc['G2'] = f'=IF(X1>0, {X}( "Other/Svc" , 0 , "Out" , "Term" , -5 , "Flag", {Q}!B8 ), 0)'
calc['G3'] = f'=IF(X1>0, {X}("Other/Svc", 0, "Out", "Term", 5.5, "Flag", FALSE), 0)'
calc['G4'] = f'={X}("Other/Svc", 0, "Out", "Term", Data!B2, "Flag", TRUE)'
calc['G5'] = f'={X}("Other/Svc", 0, "Out", "Term", Data!B2, "Flag", FALSE, "Note", "x""y")'
# Group 5: extras only, text constants.
calc['J2'] = f'={X}("Extra/Svc", 0, "Out", "Tier", {Q}!B7)'
calc['J3'] = f'={X}("Extra/Svc", 0, "Out", "Tier", "Silver", "Smoker", TRUE)'
calc['J4'] = f'={X}("Extra/Svc", 0, "Out")'
path = HERE / 'prefetch_plan.xlsx'
book.save(path)

plan = prefetch.sibling_plan_for_path(path)
observations = [
    ('Folder/Svc', 'Premium', [('term', ('float', 10.0)), ('rate', ('float', 0.05))]),
    ('Folder/Svc', 'Premium', [('term', ('int', 20)), ('rate', ('float', 0.05))]),
    ('Folder/Svc', 'Premium', [('term', ('float', 30.0)), ('rate', ('float', 0.06)), ('plan', ('str', 'Gold'))]),
    ('Folder/Svc', 'Value', [('age', ('float', 33.0)), ('band', ('float', 1.0))]),
    ('Folder/Svc', 'Value', [('age', ('float', 45.0)), ('band', ('float', 7.0))]),
    ('Other/Svc', 'Out', [('term', ('float', -5.0)), ('flag', ('bool', True))]),
    ('Other/Svc', 'Out', [('term', ('float', 12.0)), ('flag', ('bool', True))]),
    ('Extra/Svc', 'Out', [('tier', ('str', 'Gold'))]),
    ('Extra/Svc', 'Out', [('tier', ('str', 'Silver')), ('smoker', ('bool', True))]),
    ('Extra/Svc', 'Out', []),
    ('Nope/Svc', 'Out', [('term', ('float', 1.0))]),
]
vectors = []
for target, output, pairs in observations:
    observed = {name: untyped(value) for name, value in pairs}
    found = prefetch.sibling_vectors(plan, target, output, observed)
    vectors.append({'target': target, 'output': output,
                    'inputs': [[name, typed(untyped(value))] for name, value in pairs],
                    'vectors': [[[name, typed(value)] for name, value in vector.items()] for vector in found]})
expected = {
    'counts': plan.counts(),
    'groups': [{'target': group.target, 'output': group.output, 'cells': list(group.cells),
                'positions': [[name, kind, [typed(value) for value in values]]
                              for name, kind, values in group.positions]}
               for group in plan.groups],
    'vectors': vectors,
}
(HERE / 'prefetch_expected.json').write_text(json.dumps(expected, indent=1) + '\n')

# -- Brent -------------------------------------------------------------------------
FUNCTIONS = {
    'cubic': lambda x: x * x * x - 2 * x - 5,
    'square2': lambda x: x * x - 2,
    'rational': lambda x: 1 / (x + 1) - 0.3,
    'linear3': lambda x: x - 3,
    'steep': lambda x: (x - 0.123456789) * 1e6,
    'flat_then_up': lambda x: max(x - 7.5, 0.0) * 1000 - 1,
}
CASES = [
    ('cubic', 0.0, 300.0, 1e-12, 100), ('cubic', 0.0, 300.0, 1.0, 25), ('cubic', 0.0, 3.0, 1e-9, 4),
    ('square2', 0.0, 2.0, 1e-15, 200), ('rational', 0.0, 10.0, 1e-10, 50), ('linear3', 0.0, 6.0, 1e-9, 25),
    ('linear3', 3.0, 6.0, 1e-9, 25), ('square2', 2.0, 3.0, 1e-9, 25), ('square2', -2.0, 2.0, 1e-9, 25),
    ('steep', 0.0, 1.0, 1e-14, 100), ('flat_then_up', 0.0, 100.0, 1e-9, 100), ('cubic', 300.0, 0.0, 1e-12, 100),
]
brent = []
for name, lower, upper, max_change, max_iterations in CASES:
    calls = []

    def objective(x, f=FUNCTIONS[name]):
        calls.append(repr(x))
        return f(x)

    record = {'function': name, 'lower': repr(lower), 'upper': repr(upper),
              'max_change': repr(max_change), 'max_iterations': max_iterations}
    try:
        result = engine_adapter._brent_solve(objective, lower, upper, max_change, max_iterations)
        record.update(outcome='ok', value=repr(result.value), iterations=result.iterations)
    except engine_adapter._MaxIterations as exc:
        record.update(outcome='max_iterations', iterations=exc.iterations)
    except engine_adapter._NoBracket:
        record.update(outcome='no_bracket')
    except engine_adapter._ZeroSlope:
        record.update(outcome='zero_slope')
    record['calls'] = calls
    brent.append(record)


def failing_at(n):
    count = [0]

    def objective(x):
        count[0] += 1
        if count[0] >= n:
            raise engine_adapter._TargetNotNumeric('target cell is not numeric')
        return FUNCTIONS['cubic'](x)
    return objective


for n in (1, 2, 5):
    try:
        engine_adapter._brent_solve(failing_at(n), 0.0, 300.0, 1e-12, 100)
        outcome = {'outcome': 'ok'}
    except engine_adapter._TargetNotNumeric as exc:
        outcome = {'outcome': 'target_not_numeric', 'iterations': exc.iterations}
    brent.append({'function': 'cubic_failing_at', 'fail_call': n, **outcome})
(HERE / 'brent_expected.json').write_text(json.dumps(brent, indent=1) + '\n')


# -- run_solves over a fake workbook ---------------------------------------------
class FakeLiteral:
    @staticmethod
    def number(value): return ('number', value)
    @staticmethod
    def int(value): return ('int', value)
    @staticmethod
    def boolean(value): return ('bool', value)
    @staticmethod
    def text(value): return ('text', value)
    @staticmethod
    def empty(): return ('empty', None)
    @staticmethod
    def error(kind, message): return ('error', kind)


class FakeEngine:
    LiteralValue = FakeLiteral


class FakeBook:
    """Sheet S: block H1:I12; change cell B2, target cell B3 = change^3 - 2*change - 5 (+ offset)."""

    def __init__(self, labels, names, formulas, offset=0.0):
        self.cells = {}
        self.formulas = dict(formulas)
        self.names = names
        self.offset = offset
        self.log = []
        for row, (label, value) in enumerate(labels, start=1):
            self.cells[('S', row, 8)] = label
            if value is not None:
                self.cells[('S', row, 9)] = value
        self.cells[('S', 2, 2)] = 1.0
        self.evaluate_all()

    def get_named_ranges(self): return [dict(row) for row in self.names]
    def get_value(self, sheet, row, col): return self.cells.get((sheet, row, col))
    def get_formula(self, sheet, row, col): return self.formulas.get((sheet, row, col))

    def set_value(self, sheet, row, col, literal):
        kind, value = literal
        self.log.append([sheet, row, col, kind, repr(value)])
        self.cells[(sheet, row, col)] = value

    def evaluate_all(self):
        x = self.cells.get(('S', 2, 2))
        self.cells[('S', 3, 2)] = (x * x * x - 2 * x - 5 + self.offset) if isinstance(x, float) else 'bad'


BLOCK = {'name': 'Xsolve_Goal', 'scope': 'workbook', 'scope_sheet': None, 'kind': 'range', 'sheet': 'S',
         'start_row': 1, 'start_col': 8, 'end_row': 12, 'end_col': 9}
TARGET_NAME = {'name': 'GoalTarget', 'scope': 'workbook', 'scope_sheet': None, 'kind': 'cell', 'sheet': 'S',
               'start_row': 3, 'start_col': 2, 'end_row': 3, 'end_col': 2}
BASE = [('Run if', True), ('Target cell', None), ('Target value', 0.0), ('By changing', None),
        ('Solve algorithm', ' Brent '), ('Max change', 1e-12), ('Max iterations', 100), ('Lower bound', 0.0),
        ('Upper bound', 300.0), ('Solve result', None), ('Solve  Iteration', None), ('Solve target', None)]
FORMULAS = {('S', 2, 9): '=GoalTarget', ('S', 4, 9): "= 'S'!$B$2 "}


def variant(**changes):
    labels = [(label, changes.get(label, value)) for label, value in BASE]
    return labels


SCENARIOS = {
    'converges': (variant(), [BLOCK, TARGET_NAME], FORMULAS),
    'coarse': (variant(**{'Max change': 1.0, 'Max iterations': 25.0}), [BLOCK, TARGET_NAME], FORMULAS),
    'target_value_5': (variant(**{'Target value': 5}), [BLOCK, TARGET_NAME], FORMULAS),
    'run_off': (variant(**{'Run if': 0}), [BLOCK, TARGET_NAME], FORMULAS),
    'bad_algorithm': (variant(**{'Solve algorithm': 'Newton'}), [BLOCK, TARGET_NAME], FORMULAS),
    'no_bracket': (variant(**{'Lower bound': 10.0}), [BLOCK, TARGET_NAME], FORMULAS),
    'max_iterations': (variant(**{'Max iterations': 3}), [BLOCK, TARGET_NAME], FORMULAS),
    'bad_numeric': (variant(**{'Max change': -1.0}), [BLOCK, TARGET_NAME], FORMULAS),
    'bad_target_ref': (variant(), [BLOCK], {('S', 2, 9): '=Nowhere!A1', ('S', 4, 9): '=B2'}),
    'guess_bounds': ([(label, value) for label, value in variant(**{'Initial guess': 2.0})
                      if label not in ('Lower bound', 'Upper bound')] + [('Initial guess', 2.0)],
                     [BLOCK, TARGET_NAME], FORMULAS),
}
goal = []
for name, (labels, names, formulas) in SCENARIOS.items():
    fake = FakeBook(labels, names, formulas)
    record = {'scenario': name, 'labels': [[label, typed(value)] for label, value in labels],
              'names': names, 'formulas': [[s, r, c, f] for (s, r, c), f in formulas.items()]}
    try:
        results, notes, receipts = solve.run_solves(fake, FakeEngine)
        record.update(outcome='ok', notes=list(notes),
                      results={key: [repr(v['TargetValue']), repr(v['ByChangingCellValue'])]
                               for key, v in results.items()},
                      records=[{k: v for k, v in r.items() if k in ('suffix', 'root', 'iterations', 'status')}
                               for r in receipts])
        for r in record['records']:
            r['root'] = repr(r['root'])
    except solve.RequiredSolverFailure as exc:
        record.update(outcome='failed', notes=list(exc.notes),
                      records=[{k: v for k, v in r.items() if k in ('type', 'status', 'name', 'reason', 'category',
                                                                      'exception_type')} for r in exc.records])
    record['writes'] = fake.log
    goal.append(record)
(HERE / 'goal_seek_expected.json').write_text(json.dumps(goal, indent=1) + '\n')
print('ok', len(expected['groups']), 'groups;', len(brent), 'brent cases;', len(goal), 'solve scenarios')
